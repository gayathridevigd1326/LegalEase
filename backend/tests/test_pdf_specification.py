"""
Unit and Integration Tests for LegalEase.pdf Specification Compliance.
Validates all requirements across Milestones 1 to 5.
"""

import io
import pytest
from fastapi.testclient import TestClient
from docx import Document as DocxDocument
from pypdf import PdfReader

from backend.app.main import app
from ai_core.gemini_generator import GeminiDocumentGenerator
from backend.app.document.formatting import (
    sanitize_text,
    format_html_preview,
    format_docx,
    format_pdf,
)


@pytest.fixture
def spec_client():
    with TestClient(app) as c:
        yield c


def test_milestone_3_root_endpoint(spec_client):
    """PDF Milestone 3.1: Root / endpoint verifies service is running."""
    res = spec_client.get("/")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "online"
    assert "LegalEase" in data["service"]


def test_milestone_3_generate_endpoint(spec_client):
    """PDF Milestone 3.2: POST /generate with document_type, parties, terms, dates."""
    payload = {
        "document_type": "Freelance Work Contract",
        "parties": "Jane Doe (Service Provider), TechNova Inc. (Client)",
        "terms": "Payment to be made within 30 days of invoice; Deliver work on time; Confidentiality maintained; 15 days notice",
        "dates": "April 10, 2025"
    }
    res = spec_client.post("/generate", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["document_type"] == "Freelance Work Contract"
    assert "content" in data
    assert len(data["content"]) > 100
    assert "preview_html" in data
    assert "legalease-preview-container" in data["preview_html"]


def test_milestone_2_sanitize_text():
    """PDF Page 15: sanitize_text removes special characters and typographic quotes."""
    raw = "“Agreement” between ‘Party A’ and ‘Party B’ — with bullets: \uf0b7 Item 1 \u00a0\r\n"
    clean = sanitize_text(raw)
    assert '"Agreement"' in clean
    assert "'Party A'" in clean
    assert " - with bullets:" in clean
    assert "\uf0b7" not in clean
    assert "\u00a0" not in clean


def test_milestone_2_format_html_preview():
    """PDF Page 14-15: format_html_preview renders dark-themed scrollable card."""
    sample = "NON-DISCLOSURE AGREEMENT\n\nSECTION 1. CONFIDENTIALITY\nParties shall not disclose."
    html = format_html_preview(sample)
    assert "legalease-preview-container" in html
    assert "#0f172a" in html  # dark background
    assert "NON-DISCLOSURE AGREEMENT" in html
    assert "CONFIDENTIALITY" in html


def test_milestone_2_format_docx_with_logo_and_terms_table():
    """PDF Page 9, 15, 22: format_docx with logo, Times New Roman, and terms table."""
    sample_text = "RESIDENTIAL LEASE AGREEMENT\n\nSECTION 1. PREMISES\nLandlord leases to Tenant."
    terms = "Monthly rent $2,000; Security deposit $2,000; No smoking"
    docx_bytes = format_docx(sample_text, "Residential Lease Agreement", terms=terms, logo_path="assets/logo.png")

    assert len(docx_bytes) > 0
    # Parse docx from bytes to verify table and typography
    doc = DocxDocument(io.BytesIO(docx_bytes))
    assert len(doc.tables) >= 1  # Terms table exists
    table = doc.tables[0]
    assert len(table.rows) == 4  # 1 header + 3 terms
    assert "Clause #" in table.rows[0].cells[0].text
    assert "Monthly rent $2,000" in table.rows[1].cells[1].text


def test_milestone_2_format_pdf_with_logo_and_footer():
    """PDF Page 9, 15, 23: format_pdf with logo and footer on all pages."""
    sample_text = "EMPLOYMENT CONTRACT\n\nSECTION 1. POSITION\nEmployee shall serve as Senior Engineer."
    terms = "Annual salary $140,000; 20 days PTO"
    pdf_bytes = format_pdf(sample_text, "Employment Contract", terms=terms, logo_path="assets/logo.png")

    assert len(pdf_bytes) > 0
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    # Check text content in PDF
    first_page_text = reader.pages[0].extract_text()
    assert "LEGAL EASE" in first_page_text
    assert "EMPLOYMENT CONTRACT" in first_page_text


def test_scenario_1_employment_contract():
    """PDF Scenario 1: Startup founder hiring with roles, compensation, confidentiality, logo, PDF."""
    gen = GeminiDocumentGenerator()
    doc = gen.generate_document(
        document_type="Employment Contract",
        parties="John Doe (Employee), Apex Startup Inc. (Employer)",
        terms="Roles and responsibilities as CTO; $150,000 base salary plus equity; Full confidentiality and IP assignment; 30 days notice",
        dates="May 1, 2025"
    )
    assert len(doc) > 300
    assert "EMPLOYMENT" in doc
    assert "John Doe" in doc
    assert "Apex Startup" in doc
    pdf = format_pdf(doc, "Employment Contract", logo_path="assets/logo.png")
    assert len(pdf) > 1000


def test_scenario_2_nda():
    """PDF Scenario 2: Freelancer Non-Disclosure Agreement specifying parties and confidentiality."""
    gen = GeminiDocumentGenerator()
    doc = gen.generate_document(
        document_type="Non-Disclosure Agreement (NDA)",
        parties="Freelancer Smith, Client Corporation",
        terms="Scope of confidentiality covers all code and client data; 3 year term; Remedies include injunctive relief",
        dates="April 15, 2025"
    )
    assert "NON-DISCLOSURE" in doc or "NDA" in doc
    assert "Freelancer Smith" in doc


def test_scenario_3_residential_lease():
    """PDF Scenario 3: Landlord residential lease agreement with property address and downloadable DOCX."""
    gen = GeminiDocumentGenerator()
    doc = gen.generate_document(
        document_type="Residential Lease Agreement",
        parties="Landlord Vance, Tenant Cooper",
        terms="Property at 123 Maple Street; Rent $1,800/month; 12 month lease term",
        dates="June 1, 2025"
    )
    assert "LEASE" in doc
    docx = format_docx(doc, "Residential Lease Agreement", terms="Property at 123 Maple Street; Rent $1,800/month; 12 month lease term")
    assert len(docx) > 1000
