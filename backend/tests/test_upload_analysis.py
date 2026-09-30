import io
import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter


def test_upload_txt_file_and_analyze(client, auth_headers):
    sample_legal_text = (
        "EMPLOYMENT AND CONFIDENTIALITY AGREEMENT\n\n"
        "This Agreement is entered into on January 1, 2026, between Quantum Corp ('Employer') and Alice Smith ('Employee').\n\n"
        "1. Position and Compensation: Employee will serve as Director of Engineering for $160,000 annually.\n"
        "2. Confidentiality: Employee shall protect all proprietary software and trade secrets.\n"
        "3. Governing Law: This contract shall be governed by Delaware law."
    )
    file_bytes = sample_legal_text.encode("utf-8")

    res = client.post(
        "/api/uploads",
        files={"file": ("contract.txt", file_bytes, "text/plain")},
        headers=auth_headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["data"]["filename"] == "contract.txt"
    assert data["data"]["file_type"] == "txt"
    assert data["data"]["analysis"] is not None
    assert len(data["data"]["analysis"]["key_clauses"]) >= 2
    assert len(data["data"]["analysis"]["risks"]) >= 1


def test_upload_docx_file_and_analyze(client, auth_headers):
    # Create an in-memory docx file
    doc = DocxDocument()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", 0)
    doc.add_paragraph("This Non-Disclosure Agreement is executed between Disclosing Party and Receiving Party.")
    doc.add_paragraph("Receiving party covenants not to disclose confidential trade secrets for 2 years.")
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    res = client.post(
        "/api/uploads",
        files={"file": ("agreement.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=auth_headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["data"]["file_type"] == "docx"
    assert "summary" in data["data"]["analysis"]


def test_upload_pdf_file_and_analyze(client, auth_headers):
    bio = io.BytesIO()
    c = canvas.Canvas(bio, pagesize=letter)
    c.drawString(100, 750, "COMMERCIAL LEASE AGREEMENT")
    c.drawString(100, 730, "Landlord hereby leases 123 Market Street to Tenant for $3,000 per month.")
    c.drawString(100, 710, "Tenant will provide security deposit upon signature.")
    c.save()
    bio.seek(0)

    res = client.post(
        "/api/uploads",
        files={"file": ("lease.pdf", bio.getvalue(), "application/pdf")},
        headers=auth_headers
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["data"]["file_type"] == "pdf"


def test_upload_invalid_extension(client, auth_headers):
    res = client.post(
        "/api/uploads",
        files={"file": ("script.exe", b"binarycontent", "application/octet-stream")},
        headers=auth_headers
    )
    assert res.status_code == 400
    assert "Unsupported file extension" in res.json()["error"]["message"]
