import io
from datetime import datetime
from typing import Optional, Dict, Any, List
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

from backend.app.schemas.document import StructuredDocumentContent
from backend.app.document.formatting import (
    sanitize_text,
    format_html_preview,
    format_docx,
    format_pdf,
)


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas for ReportLab to compute total page numbers accurately."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(HexColor("#64748B"))

        # Running header on pages after page 1
        if self._pageNumber > 1:
            self.drawString(54, letter[1] - 36, "LEGAL EASE  |  CONFIDENTIAL LEGAL DRAFT")
            self.setStrokeColor(HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(54, letter[1] - 42, letter[0] - 54, letter[1] - 42)

        # Running footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 36, page_str)
        self.drawString(
            54, 36,
            "LegalEase AI-Drafted Document — Informational Only. Consult Qualified Legal Counsel."
        )
        self.setStrokeColor(HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 48, letter[0] - 54, 48)
        self.restoreState()


class DocumentExporter:
    """Exports structured legal documents into TXT, DOCX, and PDF formats."""

    sanitize_text = staticmethod(sanitize_text)
    format_html_preview = staticmethod(format_html_preview)
    format_docx = staticmethod(format_docx)
    format_pdf = staticmethod(format_pdf)

    @staticmethod
    def export_txt(doc_content: StructuredDocumentContent) -> bytes:
        """Export as clean, structured plain text."""
        output = io.StringIO()
        div = "=" * 76

        output.write(f"{div}\n")
        output.write(f"{doc_content.title.upper()}\n")
        output.write(f"{div}\n\n")

        if doc_content.jurisdiction:
            output.write(f"JURISDICTION: {doc_content.jurisdiction}\n")
        output.write(f"DOCUMENT TYPE: {doc_content.document_type}\n")
        output.write(f"DATE GENERATED: {datetime.now().strftime('%B %d, %Y')}\n\n")

        # Disclaimer
        if doc_content.disclaimer:
            output.write(f"LEGAL DISCLAIMER:\n{doc_content.disclaimer}\n\n")

        # Parties
        if doc_content.parties:
            output.write("-" * 76 + "\n")
            output.write("PARTIES TO THIS AGREEMENT\n")
            output.write("-" * 76 + "\n\n")
            for i, party in enumerate(doc_content.parties, 1):
                role_label = f" ({party.role})" if party.role else ""
                output.write(f"Party {i}: {party.name}{role_label}\n")
                if party.company:
                    output.write(f"  Company: {party.company}\n")
                if party.address:
                    output.write(f"  Address: {party.address}\n")
                if party.email:
                    output.write(f"  Email: {party.email}\n")
                output.write("\n")

        # Sections
        for section in doc_content.sections:
            output.write("-" * 76 + "\n")
            output.write(f"{section.heading.upper()}\n")
            output.write("-" * 76 + "\n\n")
            output.write(f"{section.content}\n\n")

        # Signatures
        output.write("-" * 76 + "\n")
        output.write("SIGNATURES & EXECUTION\n")
        output.write("-" * 76 + "\n\n")
        output.write("IN WITNESS WHEREOF, the parties hereto have executed this Agreement as of the date first above written.\n\n")

        for party in doc_content.parties:
            output.write(f"For: {party.name} ({party.role})\n")
            output.write("Signature: _____________________________________\n")
            output.write("Date:      _____________________________________\n")
            output.write("Title:     _____________________________________\n\n")

        return output.getvalue().encode("utf-8")

    @staticmethod
    def export_docx(doc_content: StructuredDocumentContent) -> bytes:
        """Export as a high-fidelity formatted DOCX with legal typography."""
        doc = DocxDocument()

        # Set 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

            # Footer
            footer = section.footer
            footer_p = footer.paragraphs[0]
            footer_p.text = "LegalEase AI-Drafted Document — Informational Only. Consult Qualified Legal Counsel."
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_run = footer_p.runs[0]
            footer_run.font.name = "Times New Roman"
            footer_run.font.size = Pt(8.5)
            footer_run.font.color.rgb = RGBColor(120, 120, 120)

        # Document Header Brand
        brand_p = doc.add_paragraph()
        brand_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        b_run = brand_p.add_run("LEGAL EASE  |  DRAFT")
        b_run.font.name = "Arial"
        b_run.font.size = Pt(8.5)
        b_run.font.bold = True
        b_run.font.color.rgb = RGBColor(79, 70, 229)

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(12)
        title_p.paragraph_format.space_after = Pt(18)
        t_run = title_p.add_run(doc_content.title.upper())
        t_run.font.name = "Times New Roman"
        t_run.font.size = Pt(18)
        t_run.font.bold = True
        t_run.font.color.rgb = RGBColor(15, 23, 42)

        # Metadata box (Jurisdiction & Date)
        meta_table = doc.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        row = meta_table.rows[0]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(3.25)
        c1.width = Inches(3.25)

        p0 = c0.paragraphs[0]
        p0.add_run(f"Jurisdiction: {doc_content.jurisdiction or 'General'}").font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(f"Date: {datetime.now().strftime('%B %d, %Y')}").font.size = Pt(9.5)

        doc.add_paragraph()  # spacer

        # Disclaimer Callout Box
        disc_p = doc.add_paragraph()
        disc_p.paragraph_format.space_before = Pt(6)
        disc_p.paragraph_format.space_after = Pt(14)
        d_run = disc_p.add_run(f"DISCLAIMER: {doc_content.disclaimer}")
        d_run.font.name = "Times New Roman"
        d_run.font.size = Pt(8.5)
        d_run.font.italic = True
        d_run.font.color.rgb = RGBColor(100, 116, 139)

        # Parties Intro
        if doc_content.parties:
            parties_intro = doc.add_paragraph()
            parties_intro.paragraph_format.space_after = Pt(8)
            p_run = parties_intro.add_run("PARTIES")
            p_run.font.name = "Times New Roman"
            p_run.font.size = Pt(12)
            p_run.font.bold = True

            parties_body = doc.add_paragraph()
            parties_body.paragraph_format.space_after = Pt(12)
            parties_body.paragraph_format.line_spacing = 1.2
            party_descs = []
            for party in doc_content.parties:
                desc = f"{party.name}"
                if party.role:
                    desc += f", hereinafter referred to as the \"{party.role}\""
                if party.company:
                    desc += f" (representing {party.company})"
                if party.address:
                    desc += f", located at {party.address}"
                party_descs.append(desc)
            parties_body.add_run("THIS AGREEMENT is entered into between " + " and ".join(party_descs) + ".").font.name = "Times New Roman"

        # Sections
        for section in doc_content.sections:
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(14)
            h_p.paragraph_format.space_after = Pt(4)
            h_p.paragraph_format.keep_with_next = True
            h_run = h_p.add_run(section.heading)
            h_run.font.name = "Times New Roman"
            h_run.font.size = Pt(12)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(15, 23, 42)

            # Content paragraphs
            paras = section.content.split("\n\n")
            for para_text in paras:
                if not para_text.strip():
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                c_run = p.add_run(para_text.strip())
                c_run.font.name = "Times New Roman"
                c_run.font.size = Pt(10.5)

        # Signatures
        doc.add_paragraph()  # spacer
        sig_head = doc.add_paragraph()
        sig_head.paragraph_format.space_before = Pt(16)
        sig_head.paragraph_format.space_after = Pt(6)
        sig_head.paragraph_format.keep_with_next = True
        s_run = sig_head.add_run("EXECUTION AND SIGNATURES")
        s_run.font.name = "Times New Roman"
        s_run.font.size = Pt(12)
        s_run.font.bold = True

        sig_intro = doc.add_paragraph()
        sig_intro.paragraph_format.space_after = Pt(12)
        sig_intro.add_run("IN WITNESS WHEREOF, the parties have caused this Agreement to be executed by their duly authorized representatives:").font.name = "Times New Roman"

        if doc_content.parties:
            num_parties = len(doc_content.parties)
            sig_table = doc.add_table(rows=1, cols=min(num_parties, 2))
            sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            row = sig_table.rows[0]

            for i, party in enumerate(doc_content.parties[:2]):
                cell = row.cells[i]
                cell.width = Inches(3.25)
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.3
                cp.add_run(f"For: {party.name} ({party.role})\n\n").bold = True
                cp.add_run("Signature: ___________________________\n\n")
                cp.add_run("Name:      ___________________________\n\n")
                cp.add_run("Title:     ___________________________\n\n")
                cp.add_run("Date:      ___________________________\n")

        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()

    @staticmethod
    def export_pdf(doc_content: StructuredDocumentContent) -> bytes:
        """Export as a publication-ready PDF using ReportLab with custom running headers and footers."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'LegalTitle',
            parent=styles['Heading1'],
            fontName='Times-Bold',
            fontSize=16,
            leading=20,
            alignment=1,  # Center
            textColor=HexColor('#0F172A'),
            spaceAfter=12
        )

        meta_style = ParagraphStyle(
            'LegalMeta',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=HexColor('#475569'),
            alignment=1,
            spaceAfter=14
        )

        disclaimer_style = ParagraphStyle(
            'LegalDisclaimer',
            parent=styles['Normal'],
            fontName='Times-Italic',
            fontSize=8,
            leading=11,
            textColor=HexColor('#64748B'),
            spaceAfter=14
        )

        heading_style = ParagraphStyle(
            'LegalHeading',
            parent=styles['Heading2'],
            fontName='Times-Bold',
            fontSize=11,
            leading=15,
            textColor=HexColor('#0F172A'),
            spaceBefore=12,
            spaceAfter=4,
            keepWithNext=True
        )

        body_style = ParagraphStyle(
            'LegalBody',
            parent=styles['Normal'],
            fontName='Times-Roman',
            fontSize=10,
            leading=14,
            textColor=HexColor('#1E293B'),
            spaceAfter=6
        )

        sig_label_style = ParagraphStyle(
            'LegalSigLabel',
            parent=styles['Normal'],
            fontName='Times-Roman',
            fontSize=9,
            leading=13,
            textColor=HexColor('#0F172A')
        )

        story = []

        # Title
        story.append(Paragraph(doc_content.title.upper(), title_style))

        # Metadata
        meta_text = f"Jurisdiction: {doc_content.jurisdiction or 'General'} &nbsp;&bull;&nbsp; Date: {datetime.now().strftime('%B %d, %Y')}"
        story.append(Paragraph(meta_text, meta_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=HexColor('#CBD5E1'), spaceAfter=10))

        # Disclaimer
        if doc_content.disclaimer:
            story.append(Paragraph(f"<b>DISCLAIMER:</b> {doc_content.disclaimer}", disclaimer_style))
            story.append(Spacer(1, 6))

        # Parties Preamble
        if doc_content.parties:
            story.append(Paragraph("PARTIES TO THIS AGREEMENT", heading_style))
            party_descs = []
            for party in doc_content.parties:
                desc = f"<b>{party.name}</b>"
                if party.role:
                    desc += f", hereinafter referred to as the <i>\"{party.role}\"</i>"
                if party.company:
                    desc += f" (representing {party.company})"
                if party.address:
                    desc += f", located at {party.address}"
                party_descs.append(desc)
            p_text = "THIS AGREEMENT is made between " + " and ".join(party_descs) + "."
            story.append(Paragraph(p_text, body_style))
            story.append(Spacer(1, 8))

        # Sections
        for section in doc_content.sections:
            story.append(Paragraph(section.heading.upper(), heading_style))
            paras = section.content.split("\n\n")
            for para in paras:
                if para.strip():
                    story.append(Paragraph(para.strip().replace("\n", "<br/>"), body_style))

        # Execution / Signature Block
        story.append(Spacer(1, 14))
        story.append(Paragraph("EXECUTION AND SIGNATURES", heading_style))
        story.append(Paragraph(
            "IN WITNESS WHEREOF, the parties hereto have executed this Agreement by their duly authorized representatives as of the date first above written.",
            body_style
        ))
        story.append(Spacer(1, 10))

        if doc_content.parties:
            sig_data = []
            col_widths = [250, 250]
            row_items = []
            for party in doc_content.parties[:2]:
                cell_text = (
                    f"<b>For: {party.name} ({party.role})</b><br/><br/>"
                    f"Signature: ___________________________<br/><br/>"
                    f"Name: _______________________________<br/><br/>"
                    f"Title: ________________________________<br/><br/>"
                    f"Date: ________________________________"
                )
                row_items.append(Paragraph(cell_text, sig_label_style))

            if len(row_items) == 1:
                row_items.append(Paragraph("", sig_label_style))

            sig_data.append(row_items)
            sig_table = Table(sig_data, colWidths=col_widths)
            sig_table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 12),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ]))
            story.append(KeepTogether(sig_table))

        doc.build(story, canvasmaker=NumberedCanvas)
        return buffer.getvalue()
